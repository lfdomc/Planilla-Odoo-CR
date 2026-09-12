"""
Auditoria Comparativa de Aguinaldo -- Odoo vs. Excel de la Empresa.

Genera un reporte Excel que muestra, lado a lado para cada empleado,
el calculo real que hace Odoo (usando el MISMO metodo centralizado
rate_helper.calc_aguinaldo_periodo() que ya usan el Wizard de
Aguinaldo, la Liquidacion, y el Simulador -- confirmado unificado)
contra el valor equivalente que reporta el Excel externo de la
empresa, marcando visualmente cualquier diferencia real -- para poder
confirmar con certeza que ambos sistemas coinciden, o encontrar con
precision donde estan las inconsistencias reales si no coinciden.
"""
import io
import base64
import logging
from datetime import date
from odoo import models, fields
from odoo.exceptions import UserError

_logger = logging.getLogger(__name__)

try:
    from openpyxl import load_workbook
except ImportError:
    load_workbook = None


class AguinaldoAuditoriaWizard(models.TransientModel):
    _name = 'planilla.aguinaldo.auditoria.wizard'
    _description = 'Auditoria Comparativa de Aguinaldo (Odoo vs Excel Empresa)'

    company_id = fields.Many2one(
        'res.company', required=True,
        default=lambda self: self.env.company,
    )
    year = fields.Integer(
        string='Año del Aguinaldo', required=True,
        default=lambda self: date.today().year,
        help='Periodo legal: 1 de diciembre del año anterior al 30 de '
             'noviembre de este año (Art. 228 CT).',
    )
    excel_file = fields.Binary(
        string='Excel de la Empresa', required=True,
        help='Cargue el Excel real de la empresa, con la hoja "Agui." '
             'que contiene el desglose quincenal de aguinaldo por '
             'empleado (columnas I Dic-Ant, II Dic-Ant, ... hasta '
             'Acumulado). Se usa unicamente para comparar -- nunca se '
             'modifica ni se usa como fuente de calculo real.',
    )
    excel_filename = fields.Char(string='Nombre de archivo')
    sheet_name = fields.Char(
        string='Nombre de la Hoja', default='Agui.',
        help='Nombre exacto de la hoja dentro del Excel que tiene el '
             'desglose de aguinaldo (normalmente "Agui.").',
    )
    hasta_columna = fields.Selection([
        ('ago', 'Hasta Agosto'), ('set', 'Hasta Septiembre'),
        ('oct', 'Hasta Octubre'), ('nov', 'Hasta Noviembre (año completo)'),
    ], string='Excel Actualizado Hasta', default='ago', required=True,
        help='Hasta que mes tiene datos reales el Excel que esta '
             'cargando -- se usa para no comparar contra columnas '
             'vacias (en 0) que el Excel aun no ha llenado.',
    )

    def _leer_excel_empresa(self):
        """
        Lee la hoja de aguinaldo del Excel de la empresa, retornando un
        dict {nombre_empleado: monto_acumulado_real} -- usando SOLO las
        columnas hasta el mes indicado en hasta_columna, para no incluir
        columnas en cero que el Excel de la empresa aun no ha llenado.
        """
        if not load_workbook:
            raise UserError('La libreria openpyxl no esta instalada en el servidor.')
        if not self.excel_file:
            raise UserError('Cargue el Excel de la empresa primero.')

        raw = base64.b64decode(self.excel_file)
        wb = load_workbook(io.BytesIO(raw), data_only=True, read_only=True)
        sheet = self.sheet_name or 'Agui.'
        if sheet not in wb.sheetnames:
            raise UserError(
                f'La hoja "{sheet}" no existe en el archivo. '
                f'Hojas disponibles: {", ".join(wb.sheetnames)}')
        ws = wb[sheet]

        # Columna 2 = I Dic-Ant, ... hasta la columna limite segun
        # hasta_columna (cada mes = 2 columnas: quincena I y II).
        limite_col = {
            'ago': 19,   # hasta II Ago
            'set': 21,   # hasta II Set
            'oct': 23,   # hasta II Oct
            'nov': 25,   # hasta II Nov
        }[self.hasta_columna]

        empresa = {}
        for row in ws.iter_rows(min_row=4, values_only=False):
            nombre_cell = row[0].value if row else None
            if not nombre_cell:
                continue
            nombre = str(nombre_cell).strip()
            total = 0.0
            for cell in row[1:limite_col]:
                v = cell.value
                if isinstance(v, (int, float)):
                    total += v
            if total:
                empresa[nombre] = round(total, 2)
        return empresa

    def action_generate(self):
        self.ensure_one()
        try:
            import xlsxwriter
        except ImportError:
            raise UserError('xlsxwriter no instalado.')

        empresa = self._leer_excel_empresa()

        # -- Calcular el equivalente real en Odoo, usando el MISMO
        # metodo centralizado (rate_helper.calc_aguinaldo_periodo) ya
        # unificado en Wizard de Aguinaldo, Liquidacion, y Simulador --
        # desde el dia siguiente al corte del Acumulado Inicial de cada
        # empleado (mismo criterio de decision que
        # calc_aguinaldo_completo usa internamente), hasta la fecha
        # limite indicada por el usuario (que puede ser un rango
        # parcial del año, ej. "hasta agosto" -- por eso esta auditoria
        # llama directo a calc_aguinaldo_periodo en vez de
        # calc_aguinaldo_completo, que siempre fuerza noviembre como
        # limite en el caso normal).
        from dateutil.relativedelta import relativedelta
        from datetime import timedelta
        meses_limite = {'ago': 8, 'set': 9, 'oct': 10, 'nov': 11}[self.hasta_columna]
        fecha_limite = date(self.year, meses_limite, 1) + relativedelta(months=1) - relativedelta(days=1)

        rh = self.env['planilla.rate.helper']
        empleados = self.env['hr.employee'].search([
            ('company_id', '=', self.company_id.id),
            ('active', '=', True),
        ])

        filas = []
        for emp in empleados:
            ag_init_amount = emp.aguinaldo_initial_amount or 0.0
            ag_init_date = emp.aguinaldo_initial_date
            dic_start = date(self.year - 1, 12, 1)

            if ag_init_amount and ag_init_date and ag_init_date >= dic_start:
                fecha_desde = ag_init_date + timedelta(days=1)
                emp_initial = ag_init_amount
            else:
                fecha_desde = dic_start
                emp_initial = 0.0

            resultado = rh.calc_aguinaldo_periodo(emp, fecha_desde, fecha_limite)
            odoo_sistema = round(resultado['total'] / 12.0, 2)
            odoo_total = round(emp_initial + odoo_sistema, 2)

            excel_total = empresa.get(emp.name)
            if excel_total is None and odoo_total == 0:
                continue  # ni Odoo ni el Excel tienen datos de este empleado

            diferencia = round(odoo_total - (excel_total or 0.0), 2)
            pct = (round(abs(diferencia) / excel_total * 100, 2)
                   if excel_total else (100.0 if odoo_total else 0.0))

            filas.append({
                'nombre': emp.name,
                'sucursal': emp.branch_id.name or '',
                'ag_init_date': ag_init_date,
                'emp_initial': emp_initial,
                'slip_count': resultado['slip_count'],
                'odoo_sistema': odoo_sistema,
                'odoo_total': odoo_total,
                'excel_total': excel_total,
                'diferencia': diferencia,
                'pct': pct,
                'en_excel': excel_total is not None,
            })

        filas.sort(key=lambda f: abs(f['diferencia']), reverse=True)

        # -- Construir el Excel de salida --------------------------------
        output = io.BytesIO()
        wb_out = xlsxwriter.Workbook(output, {'in_memory': True})
        ws_out = wb_out.add_worksheet('Auditoria Aguinaldo')

        def F(**kw):
            return wb_out.add_format(kw)

        title_fmt = F(bold=True, font_size=13, bg_color='#1F4E79', font_color='white')
        hdr_fmt = F(bold=True, bg_color='#2E4057', font_color='white',
                    border=1, text_wrap=True, align='center', valign='vcenter')
        lbl_fmt = F(align='left', border=1)
        num_fmt = F(num_format='#,##0.00', border=1)
        num_bold_fmt = F(num_format='#,##0.00', border=1, bold=True)
        ok_fmt = F(num_format='#,##0.00', border=1, bg_color='#C6EFCE', font_color='#006100')
        bad_fmt = F(num_format='#,##0.00', border=1, bg_color='#FFC7CE', font_color='#9C0006', bold=True)
        pct_ok_fmt = F(num_format='0.00"%"', border=1, bg_color='#C6EFCE', font_color='#006100')
        pct_bad_fmt = F(num_format='0.00"%"', border=1, bg_color='#FFC7CE', font_color='#9C0006', bold=True)
        no_data_fmt = F(align='center', border=1, bg_color='#FFF2CC', italic=True)

        ws_out.merge_range(0, 0, 0, 8,
            f'AUDITORIA COMPARATIVA DE AGUINALDO -- Año {self.year} -- '
            f'Odoo (metodo centralizado unificado) vs. Excel de la Empresa '
            f'(hasta {dict(self._fields["hasta_columna"].selection)[self.hasta_columna]})',
            title_fmt)
        ws_out.set_row(0, 20)

        headers = [
            'Empleado', 'Sucursal', 'Fecha Corte\nAcumulado Inicial',
            'Acumulado\nInicial (ficha)', 'Boletas\nUsadas (Odoo)',
            'Aguinaldo\nSistema (Odoo)', 'TOTAL ODOO\n(Inicial+Sistema)',
            'TOTAL EXCEL\n(Empresa)', 'Diferencia\n(Odoo - Excel)', '% Diferencia',
        ]
        for ci, h in enumerate(headers):
            ws_out.write(2, ci, h, hdr_fmt)
        anchos = [30, 14, 14, 15, 10, 16, 16, 16, 14, 12]
        for ci, w in enumerate(anchos):
            ws_out.set_column(ci, ci, w)
        ws_out.set_row(2, 30)

        row = 3
        for f in filas:
            ws_out.write(row, 0, f['nombre'], lbl_fmt)
            ws_out.write(row, 1, f['sucursal'], lbl_fmt)
            ws_out.write(row, 2, f['ag_init_date'].strftime('%d/%m/%Y') if f['ag_init_date'] else '', lbl_fmt)
            ws_out.write(row, 3, f['emp_initial'], num_fmt)
            ws_out.write(row, 4, f['slip_count'], num_fmt)
            ws_out.write(row, 5, f['odoo_sistema'], num_fmt)
            ws_out.write(row, 6, f['odoo_total'], num_bold_fmt)
            if f['en_excel']:
                ws_out.write(row, 7, f['excel_total'], num_bold_fmt)
                es_ok = abs(f['diferencia']) < 1.0
                ws_out.write(row, 8, f['diferencia'], ok_fmt if es_ok else bad_fmt)
                ws_out.write(row, 9, f['pct'] / 100, pct_ok_fmt if es_ok else pct_bad_fmt)
            else:
                ws_out.write(row, 7, 'No encontrado en Excel', no_data_fmt)
                ws_out.write(row, 8, '', no_data_fmt)
                ws_out.write(row, 9, '', no_data_fmt)
            row += 1

        row += 1
        note_fmt = F(italic=True, font_color='#666666', font_size=9)
        ws_out.merge_range(row, 0, row, 9,
            'Metodo Odoo: rate_helper.calc_aguinaldo_periodo() -- suma el salario bruto real '
            '(base_salary + horas extras + bonos + vacaciones) de cada boleta confirmada/pagada '
            'entre el dia siguiente al corte del Acumulado Inicial y la fecha limite indicada, '
            'resta el subsidio CCSS de incapacidad y suma el subsidio patronal de los dias 1-3 '
            '(Art. 79 CT), luego divide el total entre 12. Confirmado legalmente (Ley 2412, Art. 2) '
            'que este es el metodo correcto: promedio de salarios REALMENTE DEVENGADOS, no un '
            'promedio de muestra multiplicado por meses. Filas en verde: diferencia menor a 1 colon '
            '(coincidencia). Filas en rojo: diferencia real que amerita revision.',
            note_fmt)

        wb_out.close()
        xlsx_data = base64.b64encode(output.getvalue()).decode()
        filename = f'Auditoria_Aguinaldo_{self.year}.xlsx'

        attachment = self.env['ir.attachment'].create({
            'name': filename,
            'type': 'binary',
            'datas': xlsx_data,
            'mimetype': 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        })
        return {
            'type': 'ir.actions.act_url',
            'url': f'/web/content/{attachment.id}?download=true',
            'target': 'self',
        }
